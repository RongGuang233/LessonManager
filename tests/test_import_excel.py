"""Synthetic fixtures only: no private workbook contents or identities."""
import importlib.util
from pathlib import Path
import tempfile
import unittest

from openpyxl import Workbook
from openpyxl.styles import PatternFill

MODULE = Path(__file__).resolve().parents[1] / "tools" / "import_excel.py"
SPEC = importlib.util.spec_from_file_location("import_excel", MODULE)
mod = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(mod)


class ImportTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.book = Workbook()
        self.sheet = self.book.active
        for cell, text in {"B10": "虚构学生甲", "B11": "数学", "B12": "100/h"}.items():
            self.sheet[cell] = text

    def put(self, coordinate, value, completed=False):
        self.sheet[coordinate] = value
        if completed:
            self.sheet[coordinate].fill = PatternFill("solid", fgColor="FFFF00")

    def run_import(self):
        self.book.save(self.root / "2026暑假课表.xlsx")
        return mod.Importer().run(self.root)

    def test_future_without_yellow_is_cancelled_even_with_yellow_date_and_name(self):
        self.put("B10", "虚构学生甲", True)
        self.put("C9", "2099.8.1", True)
        self.put("C10", "9:00-11:00")
        data = self.run_import()
        self.assertEqual(data["courses"][0]["date"], "2099-08-01")
        self.assertEqual(data["courses"][0]["status"], "cancelled")
        self.assertIsNone(data["courses"][0]["actual_minutes"])

    def test_detail_priority_and_same_day_separate_lessons(self):
        self.put("A2", "09:00-11:00")
        self.put("A3", "18:00-20:00")
        self.put("B1", "8.1周六")
        self.put("B2", "虚构学生甲 数学", True)
        self.put("B3", "虚构学生甲 数学", True)
        self.put("C9", "8.1周六", True)
        self.put("C10", "10:00-12:00", True)
        data = self.run_import()
        self.assertEqual(len(data["courses"]), 2)
        self.assertEqual({c["start_time"] for c in data["courses"]}, {"10:00", "18:00"})
        self.assertTrue(all(not s["balance_verified"] for s in data["students"]))

    def test_no_fill_does_not_charge_and_leave_overrides_yellow(self):
        self.put("C9", "8.1周六")
        self.put("C10", "9:00-11:00")
        self.put("D9", "8.2周日（请假）", True)
        self.put("D10", "9:00-11:00", True)
        self.put("E9", "8.3周一（请假）", True)
        self.put("E10", "9:00-11:00")
        data = self.run_import()
        self.assertEqual([c["status"] for c in data["courses"]], ["cancelled", "cancelled", "cancelled"])
        self.assertTrue(all(c["actual_minutes"] is None for c in data["courses"]))
        self.assertEqual(sum(r["kind"] == "attendance_conflict" for r in data["reviews"]), 2)

    def test_payments_refund_and_carryover_are_not_double_counted(self):
        self.put("H9", "已付10节，2000元")
        self.put("H10", "剩3节")
        self.put("H11", "已退费300")
        self.put("H12", "共20h，2000元")
        self.sheet.row_dimensions[11].hidden = True
        data = self.run_import()
        self.assertEqual([(p["kind"], p["amount_cents"]) for p in data["payments"]], [("payment", 200000), ("refund", -30000)])
        self.assertTrue(all(p["date"] is None for p in data["payments"]))
        self.assertEqual(sum(r["kind"] == "carryover_not_posted" for r in data["reviews"]), 1)
        self.assertEqual(sum(r["kind"] == "money_ambiguous" for r in data["reviews"]), 1)

    def test_normalized_time_missing_date_and_stable_identity(self):
        self.put("C9", "8..22周六", True)
        self.put("C10", "8::30-10:30", True)
        self.put("D10", "20:00-22:30", True)
        data = self.run_import()
        c, missing = data["courses"]
        self.assertEqual((c["date"], c["start_time"], c["actual_minutes"]), ("2026-08-22", "08:30", 120))
        self.assertEqual((missing["date"], missing["actual_minutes"]), (None, 150))
        self.assertTrue(missing["needs_review"])
        again = self.run_import()
        self.assertEqual([c["id"] for c in data["courses"]], [c["id"] for c in again["courses"]])

    def test_repeated_columns_and_shared_students(self):
        self.put("A2", "9:00-11:00")
        self.put("B1", "8.1周六")
        self.put("C1", "8.1周六")
        self.put("B2", "虚构学生甲、虚构学生乙 数学", True)
        self.put("C2", "虚构学生甲、虚构学生乙 数学", True)
        data = self.run_import()
        self.assertEqual(len(data["students"]), 2)
        self.assertEqual(len(data["courses"]), 2)
        self.assertEqual(len({c["student_id"] for c in data["courses"]}), 2)

    def test_schedule_day_off_is_a_note_not_a_student(self):
        self.put("A2", "9:00-11:00")
        self.put("B1", "8.1周六", True)
        self.put("B2", "调休", True)
        data = self.run_import()
        self.assertEqual([s["name"] for s in data["students"]], ["虚构学生甲"])
        self.assertEqual(data["courses"], [])
        notes = [r for r in data["reviews"] if r["kind"] == "schedule_note"]
        self.assertEqual(len(notes), 1)
        self.assertIsNone(notes[0]["student_id"])
        self.assertIsNone(notes[0]["course_id"])
        self.assertIn("B2=调休", notes[0]["source"])

    def test_old_list_only_supplements_and_keeps_unknown_english(self):
        self.put("C9", "8.1周六", True)
        self.put("C10", "9:00-11:00", True)
        old = Workbook()
        old.active.append(["虚构学生甲", "2026.8.1 数学", "2024.5.1 英语"])
        old.save(self.root / "课时表.xlsx")
        (self.root / "outputs").mkdir()
        old.save(self.root / "outputs" / "课时表.xlsx")
        data = self.run_import()
        self.assertEqual(len(data["courses"]), 2)
        english = next(c for c in data["courses"] if c["subject"] == "英语")
        self.assertIsNone(english["hourly_rate_cents"])
        self.assertIsNone(english["actual_minutes"])
        self.assertTrue(english["needs_review"])

    def test_free_trial_requires_explicit_free_note(self):
        self.put("C9", "8.1周六", True)
        self.put("C10", "9:00-10:00试听", True)
        paid_trial = self.run_import()["courses"][0]
        self.assertEqual(paid_trial["hourly_rate_cents"], 10000)
        self.put("H9", "试听免费")
        data = self.run_import()
        self.assertEqual(data["courses"][0]["hourly_rate_cents"], 0)
        self.assertNotIn("zero_rate", {r["kind"] for r in data["reviews"]})

    def test_detail_subject_wins_at_same_time(self):
        self.put("A2", "09:00-11:00")
        self.put("B1", "8.1周六")
        self.put("B2", "虚构学生甲 物理", True)
        self.put("C9", "8.1周六", True)
        self.put("C10", "9:00-11:00数学", True)
        data = self.run_import()
        self.assertEqual(len(data["courses"]), 1)
        self.assertEqual(data["courses"][0]["subject"], "数学")


if __name__ == "__main__":
    unittest.main()
