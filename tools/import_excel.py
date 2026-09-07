#!/usr/bin/env python3
"""Read the seven original tuition workbooks; never modify their contents.

CLI output contains private source text. Keep it outside the repository.
Uncertain monetary notes are review items, never invented payments.
"""
from __future__ import annotations

import argparse
from collections import Counter, defaultdict
from datetime import date, datetime, timezone
from decimal import Decimal
import json
from pathlib import Path
import re
from uuid import NAMESPACE_URL, uuid5

from openpyxl import load_workbook
from openpyxl.styles.colors import COLOR_INDEX
from openpyxl.utils.cell import range_boundaries

LAYOUTS = {
    "2025暑假课表.xlsx": ("A1:AO12", "AI14:AO73", 2025),
    "2025-2026第一学期课表.xlsx": ("A1:BG13", "BE16:BK93", 2025),
    "2026寒假课表.xlsx": ("A1:U8", "N10:T47", 2026),
    "2025-2026第二学期课表.xlsx": ("A1:BO10", "AK13:AQ86", 2026),
    "2026暑假课表.xlsx": ("A1:AK6", "B9:H47", 2026),
    "2026-2027第一学期课表.xlsx": ("A1:I6", "B9:H32", 2026),
}
SUBJECTS = ("数学", "物理", "英语")
RATE = re.compile(r"(\d+(?:\.\d+)?)\s*[/／]\s*[hH小时]")
TIME = re.compile(r"(\d{1,2})[:：](\d{2})\s*[-—–~～至]\s*(\d{1,2})[:：](\d{2})")
SCHEDULE_NOTE = re.compile(r"^(?:调休|休息|休假|放假|停课|请假)(?:[（(].*[)）])?$" )


def stable_id(kind, key):
    return str(uuid5(NAMESPACE_URL, "lessonmanager/" + kind + "/" + key))


def normalized(value):
    return re.sub(r"[:：]{2,}", ":", str(value or "").strip()).replace("..", ".")


def parse_date(value, year, first_term=False):
    s = normalized(value)
    match = re.search(r"(?:(20\d{2})[.年/-])?(\d{1,2})[.月/-](\d{1,2})(?:日)?", s)
    if not match:
        return None
    y, month, day = match.groups()
    y = int(y) if y else year + (1 if first_term and int(month) < 7 else 0)
    try:
        return date(y, int(month), int(day)).isoformat()
    except ValueError:
        return None


def parse_time(value):
    s = normalized(value)
    match = TIME.search(s)
    if not match:
        # An explicit start plus duration is valid; a duration alone has no start.
        start = re.search(r"(\d{1,2})\s*点", s)
        hours = re.search(r"(\d+(?:\.\d+)?)\s*h", s, re.I)
        return (f"{int(start[1]):02}:00" if start else None,
                int(Decimal(hours[1]) * 60) if hours else None)
    sh, sm, eh, em = map(int, match.groups())
    if sh > 23 or eh > 24 or sm > 59 or em > 59 or (eh == 24 and em):
        return None, None
    duration = eh * 60 + em - sh * 60 - sm
    if duration <= 0:
        return None, None
    hours = re.search(r"(\d+(?:\.\d+)?)\s*h", s, re.I)
    return f"{sh:02}:{sm:02}", int(Decimal(hours[1]) * 60) if hours else duration


def yellow(cell):
    if cell.fill.patternType != "solid":
        return False
    color = cell.fill.fgColor
    rgb = color.rgb if color.type == "rgb" else (
        COLOR_INDEX[color.indexed] if color.type == "indexed" and color.indexed < len(COLOR_INDEX) else "")
    if not isinstance(rgb, str) or len(rgb) < 6:
        return False
    r, g, b = [int(rgb[-6:][i:i + 2], 16) for i in (0, 2, 4)]
    return r >= 220 and g >= 200 and b < 160


def source_text(filename, sheet, cells):
    return filename + " / " + sheet + " / " + "；".join(
        f"{c.coordinate}={c.value}" for c in cells if c is not None and c.value is not None)


class Importer:
    def __init__(self):
        self.data = dict(schema_version=1, students=[], courses=[], payments=[], reviews=[], periods=[], meta={})
        self.students = {}
        self.stats = {}

    def student(self, name):
        name = str(name).strip()
        if name not in self.students:
            student = dict(id=stable_id("student", name), name=name, status="active", grade="", notes="历史导入，余额尚未核对", rates={}, balance_verified=False)
            self.students[name] = student
            self.data["students"].append(student)
        return self.students[name]

    def review(self, kind, message, source, student=None, course=None):
        key = (course["id"] if course else "") + source + kind
        item = dict(id=stable_id("review", key), student_id=student["id"] if student else None,
                    course_id=course["id"] if course else None, kind=kind, message=message,
                    source=source, status="pending", resolution="")
        if not any(r["id"] == item["id"] for r in self.data["reviews"]):
            self.data["reviews"].append(item)
        if course:
            course["needs_review"] = True

    def course(self, filename, sheet, cell, student, subject, day, start, minutes, rate, source, raw, tier):
        leave = "请假" in raw or "取消" in raw
        status = "completed" if yellow(cell) and not leave else "cancelled"
        course = dict(id=stable_id("course", f"{filename}/{sheet}/{cell.coordinate}/{student['id']}"),
                      student_id=student["id"], subject=subject or "待确认", date=day,
                      start_time=start, duration_minutes=minutes if minutes else 120,
                      actual_minutes=minutes if status == "completed" else None,
                      hourly_rate_cents=rate, status=status, notes=raw, series_id=None,
                      source=source, needs_review=False)
        # Temporary extraction information never appears in the exported schema.
        course.update(_tier=tier, _file=filename, _coord=cell.coordinate, _minutes=minutes)
        if leave and yellow(cell):
            self.review("attendance_conflict", "请假备注与黄色已上标记冲突；按请假暂不收费。", source, student, course)
        return course

    def finance(self, filename, sheet, cell, student):
        s = str(cell.value or "").strip()
        if not s:
            return
        src = source_text(filename, sheet, [cell])
        if "免费" in s:
            return
        kind = "refund" if "已退" in s else "payment" if re.search("已付|已交|已缴", s) else None
        amount = None
        if kind == "refund":
            m = re.search(r"已退(?:费|款)?\s*(\d+(?:\.\d+)?)", s)
            if m:
                amount = -int(Decimal(m[1]) * 100)
        elif kind == "payment":
            # In “已付10节，3000元”, 10 is a lesson count, not money.
            m = re.search(r"(\d+(?:\.\d+)?)\s*元", s)
            if not m:
                m = re.search(r"(?:已付|已交|已缴)\s*(\d+(?:\.\d+)?)(?![\d.])(?=\s*(?:[，,；;]|$))", s)
            if m:
                amount = int(Decimal(m[1]) * 100)
        if kind and amount is not None:
            self.data["payments"].append(dict(id=stable_id("payment", f"{filename}/{sheet}/{cell.coordinate}"),
                student_id=student["id"], date=None, kind=kind, amount_cents=amount, notes=s, source=src))
            self.review("payment_date_missing", "原文明确缴费或退款，但未记录发生日期；未从课程日期推造。", src, student)
        else:
            kind = "carryover_not_posted" if re.search("剩|余|结转|还差|还欠|差.*节", s) else "money_ambiguous"
            self.review(kind, "原文为余额、课时结转或金额性质不明确的备注，未作为新增缴费累计；请核对原账。", src, student)

    def read_term(self, path, layout):
        upper_range, lower_range, year = layout
        first_term = "第一学期" in path.name
        workbook = load_workbook(path, data_only=False)
        for sheet in workbook:
            lower, upper, blocks = [], [], {}
            left, top, right, bottom = range_boundaries(lower_range)
            starts = [r - 2 for r in range(top, min(bottom, sheet.max_row) + 1)
                      if RATE.search(str(sheet.cell(r, left).value or ""))]
            for index, name_row in enumerate(starts):
                end = starts[index + 1] - 2 if index + 1 < len(starts) else bottom
                name = str(sheet.cell(name_row, left).value or "").strip()
                if not name:
                    continue
                student = self.student(name)
                subjects = [x for x in SUBJECTS if x in str(sheet.cell(name_row + 1, left).value or "")]
                rate = int(Decimal(RATE.search(str(sheet.cell(name_row + 2, left).value))[1]) * 100)
                student["rates"].update({s: rate for s in subjects})
                blocks[name] = (student, subjects, rate)
                notes = [sheet.cell(r, right) for r in range(name_row - 1, end + 1) if str(sheet.cell(r, right).value or "").strip()]
                free_trial = any("免费" in str(c.value) and "试听" in str(c.value) for c in notes)
                for cell in notes:
                    self.finance(path.name, sheet.title, cell, student)
                for col in range(left + 1, right):
                    date_cell = None
                    for row in range(name_row - 1, end + 1):
                        cell = sheet.cell(row, col)
                        value = str(cell.value or "").strip()
                        if not value:
                            date_cell = None
                            continue
                        day = parse_date(value, year, first_term)
                        if day and not TIME.search(normalized(value)):
                            date_cell = cell
                            continue
                        start, minutes = parse_time(value)
                        if not (start or minutes or "合并" in value):
                            self.review("unparsed_detail", "个人明细非空单元格无法解析为课程，保留原文待核对。", source_text(path.name, sheet.title, [cell]), student)
                            continue
                        day = parse_date(date_cell.value, year, first_term) if date_cell else None
                        subject = next((s for s in SUBJECTS if s in value), subjects[0] if len(subjects) == 1 else None)
                        src = source_text(path.name, sheet.title, [sheet.cell(name_row, left), sheet.cell(name_row + 2, left), date_cell, cell])
                        raw = "；".join(str(c.value) for c in [date_cell, cell] if c)
                        course_rate = 0 if free_trial and ("试听" in value or row == name_row) else rate
                        if course_rate == 0 and free_trial:
                            src += "；" + source_text(path.name, sheet.title, [c for c in notes if "免费" in str(c.value)])
                        c = self.course(path.name, sheet.title, cell, student, subject, day, start, minutes, course_rate, src, raw, "detail")
                        if date_cell and yellow(date_cell) and not yellow(cell) and ("请假" in raw or "取消" in raw):
                            self.review("attendance_conflict", "日期黄色标记与请假备注冲突；按请假暂不收费。", src, student, c)
                        if "合并" in value:
                            self.review("merged_lesson", "备注说明与其他日期合并，无法确认独立课时；实际时长留空。", src, student, c)
                        lower.append(c)
            # Read every cell, including hidden rows/columns. Time labels may restart mid-table.
            _, _, upper_right, upper_bottom = range_boundaries(upper_range)
            for row in range(2, min(upper_bottom, sheet.max_row) + 1):
                time_cell = None
                for col in range(1, min(upper_right, sheet.max_column) + 1):
                    cell = sheet.cell(row, col)
                    value = str(cell.value or "").strip()
                    if not value:
                        continue
                    if TIME.fullmatch(normalized(value)):
                        time_cell = cell
                        continue
                    date_cell = sheet.cell(1, col)
                    day = parse_date(date_cell.value, year, first_term)
                    if not day:
                        continue
                    if SCHEDULE_NOTE.fullmatch(value):
                        self.review("schedule_note", "排班说明，保留原文；未作为学生或课程导入。",
                                    source_text(path.name, sheet.title, [date_cell, time_cell, cell]))
                        continue
                    name_part = re.split(r"\s|数学|物理|英语|试听|\d", value)[0].strip()
                    found = [n for n in re.split(r"[/、，,]", name_part) if n]
                    if not found:
                        self.review("unparsed_schedule", "上方排班无法识别学生，保留原文待核对。", source_text(path.name, sheet.title, [date_cell, cell]))
                    for name in found:
                        student = self.student(name)
                        _, subjects, rate = blocks.get(name, (student, [], None))
                        subject = next((s for s in SUBJECTS if s in value), subjects[0] if len(subjects) == 1 else None)
                        start, minutes = parse_time(value)
                        if start is None:
                            start, inherited = parse_time(time_cell.value if time_cell else "")
                            minutes = minutes or inherited
                        src = source_text(path.name, sheet.title, [date_cell, time_cell, cell])
                        upper.append(self.course(path.name, sheet.title, cell, student, subject, day, start, minutes, rate, src, value, "schedule"))
            # Identical repeated schedule columns are one source, not extra lessons.
            upper = self.exact_dedup(upper)
            lower = self.exact_dedup(lower)
            self.merge_upper(lower, upper, blocks)
            self.data["courses"].extend(lower)
            dates = [c["date"] for c in lower if c["date"]]
            if dates:
                self.data["periods"].append(dict(id=stable_id("period", path.name + sheet.title), name=path.stem, start=min(dates), end=max(dates)))
        workbook.close()

    @staticmethod
    def exact_dedup(courses):
        seen, result = {}, []
        for c in courses:
            key = (c["student_id"], c["date"], c["subject"], c["start_time"], c["_minutes"])
            # Unknown dates/times do not establish duplicate identity.
            if c["date"] and c["start_time"] and key in seen:
                old = seen[key]
                if old["status"] == c["status"]:
                    old["source"] += "；重复来源：" + c["source"]
                    continue
            seen[key] = c
            result.append(c)
        return result

    def merge_upper(self, lower, upper, blocks):
        groups = defaultdict(list)
        for c in lower:
            groups[c["student_id"], c["date"]].append(c)
        used = set()
        # Closest time pairing preserves multiple lessons on the same date.
        pairs = []
        def clock(t):
            return int(t[:2]) * 60 + int(t[3:]) if t else 0
        for ui, u in enumerate(upper):
            for l in groups[u["student_id"], u["date"]]:
                compatible = l["subject"] == u["subject"] or "待确认" in (l["subject"], u["subject"])
                if compatible or (l["start_time"] and l["start_time"] == u["start_time"]):
                    penalty = 0 if compatible else 1440
                    pairs.append((penalty + abs(clock(l["start_time"]) - clock(u["start_time"])), ui, l))
        matched = set()
        for _, ui, l in sorted(pairs, key=lambda p: p[0]):
            if ui in matched or l["id"] in used:
                continue
            u = upper[ui]
            used.add(l["id"])
            matched.add(ui)
            if l["subject"] == "待确认" and u["subject"] != "待确认":
                l["subject"] = u["subject"]
            l["source"] += "；上方对照：" + u["source"]
        for ui, u in enumerate(upper):
            if ui in matched:
                continue
            # Confirmed source correction: this trial is September 2 in personal detail.
            if u["_file"] == "2026-2027第一学期课表.xlsx" and u["_coord"] == "E3" and "试听" in u["notes"]:
                trials = [c for c in lower if c["student_id"] == u["student_id"] and c["date"] == "2026-09-02"]
                if len(trials) == 1:
                    trials[0]["source"] += "；日期按下方明细修正，上方原文：" + u["source"]
                    continue
            lower.append(u)
            student = next(s for s in self.data["students"] if s["id"] == u["student_id"])
            if u["status"] == "completed":
                self.review("schedule_only", "仅上方排班有此已上记录，个人明细未匹配；保留课程并待核对扣费。", u["source"], student, u)

    def read_legacy(self, path):
        workbook = load_workbook(path)
        existing = defaultdict(list)
        for c in self.data["courses"]:
            existing[c["student_id"], c["date"], c["subject"]].append(c)
        for sheet in workbook:
            for row in sheet:
                if not row[0].value:
                    continue
                student = self.student(row[0].value)
                for cell in row[1:]:
                    value = str(cell.value or "").strip()
                    if not value:
                        continue
                    day = parse_date(value, 2025)
                    subject = next((s for s in SUBJECTS if s in value), None)
                    src = source_text(path.name, sheet.title, [row[0], cell])
                    if cell.coordinate == "G4" and day == "2028-07-28" and subject == "物理":
                        candidates = existing[student["id"], "2025-07-28", "物理"]
                        if any("AM64=" in c["source"] and "2025暑假" in c["source"] for c in candidates):
                            day = "2025-07-28"
                            src += "；年份由2025暑假同学生AM62/AM64交叉证据修正为2025"
                    matches = existing[student["id"], day, subject or "待确认"]
                    if matches:
                        # The old list lacks times; it cannot establish an extra same-day lesson.
                        matches[0]["source"] += "；低优先级对照：" + src
                        continue
                    c = self.course(path.name, sheet.title, cell, student, subject, day, None, None, None, src, value, "legacy")
                    self.data["courses"].append(c)
                    existing[student["id"], day, c["subject"]].append(c)
                    self.review("legacy_incomplete", "低优先级旧表只有日期/科目，缺少价格、时长及可核验扣费明细；未据此补造金额。", src, student, c)
        workbook.close()

    def run(self, source_dir):
        source_dir = Path(source_dir)
        for filename, layout in LAYOUTS.items():
            path = source_dir / filename
            if path.is_file():
                before = len(self.data["courses"])
                self.read_term(path, layout)
                self.stats[filename] = {"courses": len(self.data["courses"]) - before}
        legacy = source_dir / "课时表.xlsx"
        if legacy.is_file():
            before = len(self.data["courses"])
            self.read_legacy(legacy)
            self.stats[legacy.name] = {"courses": len(self.data["courses"]) - before}
        if not self.stats:
            raise ValueError("未找到受支持的原始工作簿")
        student_ids = {s["id"]: s for s in self.data["students"]}
        for course in self.data["courses"]:
            missing = []
            if not course["date"]:
                missing.append("日期")
            if course["subject"] == "待确认":
                missing.append("科目")
            if course["_minutes"] is None:
                missing.append("实际时长")
            if course["hourly_rate_cents"] is None:
                missing.append("历史单价")
            if missing:
                message = "缺少" + "、".join(missing) + "，保留空值等待核对。"
                if course["_minutes"] is None:
                    message += "排期时长120分钟仅为界面占位，非实际课时。"
                self.review("missing_course_fields", message, course["source"], student_ids[course["student_id"]], course)
            if course["hourly_rate_cents"] == 0 and "免费" not in course["source"]:
                # Zero remains exactly as recorded; it is not an inferred free lesson.
                self.review("zero_rate", "原表单价为0/h，保留原值，需确认是否确实免费。", course["source"], student_ids[course["student_id"]], course)
            if course["_file"] == "2025-2026第二学期课表.xlsx" and course["_coord"] == "AN36":
                self.review("time_typo", "原文21:00-23:01可能有笔误，保留121分钟并待确认。", course["source"], student_ids[course["student_id"]], course)
            for key in list(course):
                if key.startswith("_"):
                    del course[key]
        # Dropped duplicate candidates must not leave dangling course reviews.
        course_ids = {c["id"] for c in self.data["courses"]}
        self.data["reviews"] = [r for r in self.data["reviews"] if not r["course_id"] or r["course_id"] in course_ids]
        self.data["meta"] = dict(imported_at=datetime.now(timezone.utc).isoformat(), source_summary=self.stats)
        return self.data


def summary(data):
    return {"counts": {key: len(data[key]) for key in ("students", "courses", "payments", "reviews", "periods")},
            "course_statuses": dict(Counter(c["status"] for c in data["courses"])),
            "review_kinds": dict(Counter(r["kind"] for r in data["reviews"])),
            "courses_needing_review": sum(c["needs_review"] for c in data["courses"]),
            "english_courses": sum(c["subject"] == "英语" for c in data["courses"]),
            "source_summary": data["meta"]["source_summary"],
            "balance_limitation": "全体余额未核对；缺失付款/退款日期、性质不明资金、跨期余额及未知历史课时/单价均需原账核对。"}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source-dir", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    data = Importer().run(args.source_dir)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    anonymized = summary(data)
    args.output.with_suffix(".summary.json").write_text(json.dumps(anonymized, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(anonymized, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
